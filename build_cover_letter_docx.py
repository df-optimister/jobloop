"""
커버레터 평문(문단 사이 빈 줄)을 .docx로 변환한다. CV와 같은 연락처 헤더(LinkedIn/GitHub 하이퍼링크 포함)를 얹는다.

실행:
  python build_cover_letter_docx.py <letter.txt> <candidate_name> <headline> <location> <contact_line> <out_dir> [out_name]
  contact_line 예: "you@example.com | linkedin.com/in/your-name | github.com/your-handle"
"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

import docx
from docx.shared import Pt, Inches

from build_docx import add_hyperlink


def build(letter_path: str, name: str, headline: str, location: str, contact_line: str,
          out_dir: str, out_name: str | None = None):
    text = Path(letter_path).read_text(encoding="utf-8").strip()

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.8)

    name_p = doc.add_paragraph()
    name_run = name_p.add_run(name)
    name_run.bold = True
    name_run.font.size = Pt(14)
    name_p.paragraph_format.space_after = Pt(1)

    if headline:
        hp = doc.add_paragraph(headline)
        hp.paragraph_format.space_after = Pt(1)
        hp.runs[0].font.size = Pt(10)

    if location:
        lp = doc.add_paragraph(location)
        lp.paragraph_format.space_after = Pt(1)
        lp.runs[0].font.size = Pt(10)

    cp = doc.add_paragraph()
    cp.paragraph_format.space_after = Pt(12)
    parts = [seg.strip() for seg in contact_line.split("|")]
    for i, seg in enumerate(parts):
        if i > 0:
            cp.add_run(" | ").font.size = Pt(10)
        low = seg.lower()
        if low.startswith(("linkedin.com", "github.com")):
            add_hyperlink(cp, "https://" + seg, seg, size=10, underline=False)
        else:
            cp.add_run(seg).font.size = Pt(10)

    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(10)
        for r in p.runs:
            r.font.size = Pt(11)

    out_path = Path(out_dir) / f"{out_name or Path(letter_path).stem}.docx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


if __name__ == "__main__":
    letter_path, name, headline, location, contact_line, out_dir = sys.argv[1:7]
    out_name = sys.argv[7] if len(sys.argv) > 7 else None
    p = build(letter_path, name, headline, location, contact_line, out_dir, out_name)
    print(f"saved: {p}")

    try:
        from get_page_count import page_count
        n = page_count(str(p))
        print(f"pages: {n}")
    except ImportError:
        pass
