"""
자기소개서/자유서술 문항 답변을 하나의 .docx로 묶는다: 문항마다 [번호. 질문 → 글자수(실제/한도) → 답변] 순서.
글자수는 공백·줄바꿈 포함(한국 채용 사이트 관례). 한도를 넘으면 글자수 줄을 빨간색으로 표시하고 경고를 출력한다.

입력 JSON 형식:
{
  "title": "OO기업 2026 하반기 신입 — OO 직무 자기소개서",
  "items": [
    {"question": "문항 전문", "limit": 500, "unit": "자", "answer_file": "runs/.../q1.txt"},
    {"question": "...", "limit": 250, "unit": "words", "answer": "직접 넣은 답변"}
  ],
  "overseas": [
    {"org": "Example University", "country": "영국", "period": "2018.09 ~ 2022.07", "purpose": "교육",
     "limit": 200, "answer_file": "runs/.../overseas_university.txt"}
  ]
}
unit은 "자"(공백 포함 글자수, 기본값) 또는 "words"(영문 단어수).
overseas(선택)가 있으면 문항 뒤에 "해외경험" 목차를 따로 만들어 [기관 | 국가 | 기간 | 목적 → 글자수 → 내용] 순서로 넣는다.

실행:
  python build_answers_docx.py <items.json> <output.docx>
"""
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def count(text: str, unit: str) -> int:
    return len(text.split()) if unit == "words" else len(text)


def set_east_asian_font(run, name="맑은 고딕"):
    """python-docx의 font.name은 한글(eastAsia) 글꼴을 바꾸지 않아 rFonts를 직접 지정한다."""
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)


def build(spec_path: str, out_path: str):
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    tr = title.add_run(spec["title"])
    tr.bold = True
    tr.font.size = Pt(14)
    set_east_asian_font(tr)

    over = []
    for i, item in enumerate(spec["items"], 1):
        add_entry(doc, f"{i}. {item['question']}", item, f"문항 {i}", over)

    if spec.get("overseas"):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(22)
        hr = h.add_run("해외경험")
        hr.bold = True
        hr.font.size = Pt(13)
        set_east_asian_font(hr)
        for i, item in enumerate(spec["overseas"], 1):
            head = " | ".join(x for x in (item.get("org"), item.get("country"), item.get("period"), item.get("purpose")) if x)
            add_entry(doc, f"{i}) {head}", item, f"해외경험 {i}", over)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"saved: {out_path}")
    for name, n, limit, unit in over:
        print(f"[WARN] {name}: {n}{unit} > 한도 {limit}{unit}")


def add_entry(doc, heading: str, item: dict, name: str, over: list):
    """[제목 → 글자수(실제/한도) → 답변] 한 덩어리. 한도 초과는 over에 (name, n, limit, unit)으로 쌓는다."""
    answer = item.get("answer") or Path(item["answer_file"]).read_text(encoding="utf-8")
    answer = answer.strip()
    unit = item.get("unit", "자")
    n, limit = count(answer, unit), item.get("limit")

    q = doc.add_paragraph()
    q.paragraph_format.space_before = Pt(14)
    q.paragraph_format.space_after = Pt(2)
    qr = q.add_run(heading)
    qr.bold = True
    qr.font.size = Pt(11)
    qr.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    set_east_asian_font(qr)

    c = doc.add_paragraph()
    c.paragraph_format.space_after = Pt(6)
    label = f"{n} / {limit}{unit}" if limit else f"{n}{unit}"
    if limit and n > limit:
        label += f"  (한도 {n - limit}{unit} 초과)"
        over.append((name, n, limit, unit))
    cr = c.add_run(f"글자수: {label}" if unit == "자" else f"Word count: {label}")
    cr.font.size = Pt(9)
    cr.italic = True
    cr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00) if limit and n > limit else RGBColor(0x60, 0x60, 0x60)
    set_east_asian_font(cr)

    for para in answer.split("\n"):
        if not para.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(para.strip())
        set_east_asian_font(r)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
